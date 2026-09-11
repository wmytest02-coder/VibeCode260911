import sys
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QPushButton, QTableWidget, QTableWidgetItem,
    QSpinBox, QMessageBox, QHeaderView, QGroupBox, QFrame
)
from PyQt6.QtCore import Qt, QSize
from PyQt6.QtGui import QFont, QColor, QLinearGradient, QPalette
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from MyProducts import ProductDatabase
from datetime import datetime


class ProductsGUI(QMainWindow):
    """PyQt6을 사용한 제품 관리 GUI"""
    
    def __init__(self):
        super().__init__()
        self.db = ProductDatabase("products.db")
        self.init_ui()
        self.load_products()
    
    def init_ui(self):
        """UI 초기화"""
        # 메인 윈도우 설정
        self.setWindowTitle("💼 제품 관리 시스템")
        self.setGeometry(100, 100, 1000, 700)
        self.setMinimumSize(800, 500)
        
        # 화려한 스타일시트
        self.setStyleSheet("""
            QMainWindow {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                                           stop:0 #f5f7fa, stop:1 #c3cfe2);
            }
            QGroupBox {
                color: #333333;
                border: 2px solid #0066cc;
                border-radius: 8px;
                margin-top: 12px;
                padding-top: 12px;
                background-color: rgba(255, 255, 255, 0.95);
                font-weight: bold;
                font-size: 12px;
            }
            QGroupBox::title {
                subcontrol-origin: margin;
                left: 10px;
                padding: 0 3px 0 3px;
                color: #0066cc;
            }
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #4CAF50, stop: #2E7D32);
                color: white;
                border: none;
                padding: 10px 15px;
                border-radius: 6px;
                font-weight: bold;
                font-size: 11px;
                border: 1px solid #1B5E20;
                spacing: 5px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #66BB6A, stop: #388E3C);
                border: 1px solid #2E7D32;
            }
            QPushButton:pressed {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #2E7D32, stop: #1B5E20);
            }
            QLineEdit, QSpinBox {
                padding: 10px;
                border: 2px solid #0066cc;
                border-radius: 6px;
                font-size: 11px;
                background-color: white;
                selection-background-color: #0066cc;
                color: #333333;
            }
            QLineEdit:focus, QSpinBox:focus {
                border: 2px solid #0099ff;
                background-color: #f0f8ff;
            }
            QLabel {
                font-size: 12px;
                color: #333333;
                font-weight: 500;
            }
            QTableWidget {
                background-color: white;
                alternate-background-color: #e8f4f8;
                border: 2px solid #0066cc;
                border-radius: 6px;
                gridline-color: #d0d0d0;
            }
            QTableWidget::item {
                padding: 8px;
                border-bottom: 1px solid #e0e0e0;
            }
            QTableWidget::item:selected {
                background-color: #0066cc;
                color: white;
            }
            QHeaderView::section {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #0066cc, stop: #0052a3);
                color: white;
                padding: 8px;
                border: none;
                font-weight: bold;
            }
            QMessageBox QLabel {
                color: #333333;
            }
            QMessageBox QPushButton {
                min-width: 60px;
            }
        """)
        
        # 중앙 위젯
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)
        main_layout.setSpacing(15)
        main_layout.setContentsMargins(20, 20, 20, 20)
        
        # 제목
        title_label = QLabel("💼 제품 관리 시스템")
        title_font = QFont()
        title_font.setPointSize(18)
        title_font.setBold(True)
        title_label.setFont(title_font)
        title_label.setStyleSheet("""
            QLabel {
                color: #0066cc;
                padding: 10px;
                background-color: rgba(255, 255, 255, 0.8);
                border-radius: 8px;
                border-left: 5px solid #0066cc;
            }
        """)
        main_layout.addWidget(title_label)
        
        # ========== 입력 영역 ==========
        input_group = QGroupBox("📝 제품 정보 입력")
        input_layout = QHBoxLayout(input_group)
        input_layout.setSpacing(12)
        
        # 제품명 입력
        input_layout.addWidget(QLabel("제품명:"))
        self.product_name_input = QLineEdit()
        self.product_name_input.setPlaceholderText("제품명을 입력하세요")
        self.product_name_input.setMaximumWidth(250)
        input_layout.addWidget(self.product_name_input)
        
        # 가격 입력
        input_layout.addWidget(QLabel("가격:"))
        self.product_price_input = QSpinBox()
        self.product_price_input.setRange(0, 10000000)
        self.product_price_input.setMaximumWidth(150)
        self.product_price_input.setSuffix(" 원")
        input_layout.addWidget(self.product_price_input)
        
        # 추가 버튼
        self.add_button = QPushButton("➕ 추가")
        self.add_button.setMaximumWidth(110)
        self.add_button.clicked.connect(self.add_product)
        input_layout.addWidget(self.add_button)
        
        input_layout.addStretch()
        main_layout.addWidget(input_group)
        
        # ========== 검색 영역 ==========
        search_group = QGroupBox("🔍 제품 검색")
        search_layout = QHBoxLayout(search_group)
        search_layout.setSpacing(12)
        
        search_layout.addWidget(QLabel("검색어:"))
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("제품명으로 검색...")
        self.search_input.setMaximumWidth(350)
        self.search_input.textChanged.connect(self.search_products)
        search_layout.addWidget(self.search_input)
        
        # 전체 보기 버튼
        self.show_all_button = QPushButton("🔄 전체 보기")
        self.show_all_button.setMaximumWidth(120)
        self.show_all_button.clicked.connect(self.load_products)
        search_layout.addWidget(self.show_all_button)
        
        search_layout.addStretch()
        main_layout.addWidget(search_group)
        
        # ========== 테이블 영역 ==========
        table_group = QGroupBox("📊 제품 목록")
        table_layout = QVBoxLayout(table_group)
        
        self.table = QTableWidget()
        self.table.setColumnCount(3)
        self.table.setHorizontalHeaderLabels(["ID", "제품명", "가격"])
        self.table.setAlternatingRowColors(True)
        self.table.setRowHeight(25)
        
        # 컬럼 크기 조정
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        
        # 헤더 스타일
        for i in range(3):
            item = self.table.horizontalHeaderItem(i)
            if item:
                item.setFont(QFont("Arial", 12, QFont.Weight.Bold))
                item.setForeground(QColor("white"))
        
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QTableWidget.SelectionMode.SingleSelection)
        table_layout.addWidget(self.table)
        main_layout.addWidget(table_group, 1)
        
        # ========== 버튼 영역 ==========
        button_group = QGroupBox("⚙️ 작업")
        button_layout = QHBoxLayout(button_group)
        button_layout.setSpacing(10)
        
        # 수정 버튼
        self.edit_button = QPushButton("✏️ 수정")
        self.edit_button.setMaximumWidth(110)
        self.edit_button.clicked.connect(self.edit_product)
        button_layout.addWidget(self.edit_button)
        
        # 삭제 버튼
        self.delete_button = QPushButton("🗑️ 삭제")
        self.delete_button.setMaximumWidth(110)
        self.delete_button.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #f44336, stop: #c62828);
                color: white;
                border: none;
                padding: 10px 15px;
                border-radius: 6px;
                font-weight: bold;
                font-size: 11px;
                border: 1px solid #b71c1c;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #ef5350, stop: #d32f2f);
            }
        """)
        self.delete_button.clicked.connect(self.delete_product)
        button_layout.addWidget(self.delete_button)
        
        button_layout.addStretch()
        
        # 엑셀 저장 버튼
        self.export_button = QPushButton("📥 엑셀로 저장")
        self.export_button.setMaximumWidth(140)
        self.export_button.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #2196F3, stop: #1565c0);
                color: white;
                border: none;
                padding: 10px 15px;
                border-radius: 6px;
                font-weight: bold;
                font-size: 11px;
                border: 1px solid #0d47a1;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #42a5f5, stop: #1976d2);
            }
        """)
        self.export_button.clicked.connect(self.export_to_excel)
        button_layout.addWidget(self.export_button)
        
        # 닫기 버튼
        self.close_button = QPushButton("❌ 닫기")
        self.close_button.setMaximumWidth(110)
        self.close_button.setStyleSheet("""
            QPushButton {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #757575, stop: #424242);
                color: white;
                border: none;
                padding: 10px 15px;
                border-radius: 6px;
                font-weight: bold;
                font-size: 11px;
                border: 1px solid #212121;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                                           start: #9e9e9e, stop: #616161);
            }
        """)
        self.close_button.clicked.connect(self.close)
        button_layout.addWidget(self.close_button)
        
        main_layout.addWidget(button_group)
    
    def add_product(self):
        """제품 추가"""
        product_name = self.product_name_input.text().strip()
        product_price = self.product_price_input.value()
        
        if not product_name:
            msg_box = QMessageBox(self)
            msg_box.setWindowTitle("⚠️ 경고")
            msg_box.setText("제품명을 입력해주세요.")
            msg_box.setIcon(QMessageBox.Icon.Warning)
            msg_box.exec()
            self.product_name_input.setFocus()
            return
        
        if product_price <= 0:
            msg_box = QMessageBox(self)
            msg_box.setWindowTitle("⚠️ 경고")
            msg_box.setText("가격은 0보다 커야 합니다.")
            msg_box.setIcon(QMessageBox.Icon.Warning)
            msg_box.exec()
            self.product_price_input.setFocus()
            return
        
        # 데이터베이스에 추가
        self.db.insert_product(product_name, product_price)
        
        # UI 초기화
        self.product_name_input.clear()
        self.product_price_input.setValue(0)
        
        # 테이블 새로고침
        self.load_products()
        
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("✅ 성공")
        msg_box.setText("제품이 추가되었습니다.")
        msg_box.setIcon(QMessageBox.Icon.Information)
        msg_box.exec()
    
    def load_products(self):
        """모든 제품 로드"""
        self.search_input.clear()
        products = self.db.get_all_products()
        self.display_products(products)
    
    def search_products(self):
        """제품 검색"""
        search_text = self.search_input.text().strip()
        
        if not search_text:
            self.load_products()
            return
        
        products = self.db.search_by_name(search_text)
        self.display_products(products)
    
    def display_products(self, products):
        """테이블에 제품 표시"""
        self.table.setRowCount(0)
        
        for product in products:
            row_position = self.table.rowCount()
            self.table.insertRow(row_position)
            
            # ID
            id_item = QTableWidgetItem(str(product[0]))
            id_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            id_item.setFont(QFont("Arial", 10, QFont.Weight.Bold))
            id_item.setForeground(QColor("#0066cc"))
            self.table.setItem(row_position, 0, id_item)
            
            # 제품명
            name_item = QTableWidgetItem(product[1])
            name_item.setFont(QFont("Arial", 10))
            name_item.setForeground(QColor("#333333"))
            self.table.setItem(row_position, 1, name_item)
            
            # 가격 (천의 자리 콤마 포함)
            price_item = QTableWidgetItem(f"{product[2]:,}")
            price_item.setTextAlignment(Qt.AlignmentFlag.AlignRight)
            price_item.setFont(QFont("Arial", 10, QFont.Weight.Bold))
            price_item.setForeground(QColor("#d32f2f"))
            self.table.setItem(row_position, 2, price_item)
    
    def edit_product(self):
        """제품 수정"""
        selected_row = self.table.currentRow()
        
        if selected_row < 0:
            msg_box = QMessageBox(self)
            msg_box.setWindowTitle("⚠️ 경고")
            msg_box.setText("수정할 제품을 선택해주세요.")
            msg_box.setIcon(QMessageBox.Icon.Warning)
            msg_box.exec()
            return
        
        product_id = int(self.table.item(selected_row, 0).text())
        product_name = self.product_name_input.text().strip()
        product_price = self.product_price_input.value()
        
        if not product_name and product_price == 0:
            msg_box = QMessageBox(self)
            msg_box.setWindowTitle("⚠️ 경고")
            msg_box.setText("변경할 항목을 입력해주세요.")
            msg_box.setIcon(QMessageBox.Icon.Warning)
            msg_box.exec()
            return
        
        # 업데이트 파라미터 준비
        update_name = product_name if product_name else None
        update_price = product_price if product_price > 0 else None
        
        # 데이터베이스 업데이트
        self.db.update_product(product_id, update_name, update_price)
        
        # UI 초기화
        self.product_name_input.clear()
        self.product_price_input.setValue(0)
        
        # 테이블 새로고침
        self.load_products()
        
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("✅ 성공")
        msg_box.setText("제품이 수정되었습니다.")
        msg_box.setIcon(QMessageBox.Icon.Information)
        msg_box.exec()
    
    def delete_product(self):
        """제품 삭제"""
        selected_row = self.table.currentRow()
        
        if selected_row < 0:
            msg_box = QMessageBox(self)
            msg_box.setWindowTitle("⚠️ 경고")
            msg_box.setText("삭제할 제품을 선택해주세요.")
            msg_box.setIcon(QMessageBox.Icon.Warning)
            msg_box.exec()
            return
        
        product_id = int(self.table.item(selected_row, 0).text())
        product_name = self.table.item(selected_row, 1).text()
        
        # 확인 대화상자
        msg_box = QMessageBox(self)
        msg_box.setWindowTitle("🗑️ 삭제 확인")
        msg_box.setText(f"'{product_name}'을(를) 삭제하시겠습니까?")
        msg_box.setIcon(QMessageBox.Icon.Question)
        msg_box.setStandardButtons(QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        msg_box.setDefaultButton(QMessageBox.StandardButton.No)
        
        reply = msg_box.exec()
        
        if reply == QMessageBox.StandardButton.Yes:
            self.db.delete_product(product_id)
            self.load_products()
            
            success_box = QMessageBox(self)
            success_box.setWindowTitle("✅ 성공")
            success_box.setText("제품이 삭제되었습니다.")
            success_box.setIcon(QMessageBox.Icon.Information)
            success_box.exec()
    
    def export_to_excel(self):
        """엑셀 파일로 저장"""
        products = self.db.get_all_products()
        
        if not products:
            msg_box = QMessageBox(self)
            msg_box.setWindowTitle("⚠️ 경고")
            msg_box.setText("저장할 데이터가 없습니다.")
            msg_box.setIcon(QMessageBox.Icon.Warning)
            msg_box.exec()
            return
        
        try:
            # 워크북 및 시트 생성
            wb = Workbook()
            ws = wb.active
            ws.title = "Products"
            
            # 헤더 행 스타일 (더 화려한 색상)
            header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="0066cc", end_color="0066cc", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # 헤더 행 작성
            headers = ["ID", "제품명", "가격"]
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # 데이터 행 작성
            data_alignment_center = Alignment(horizontal="center")
            data_alignment_right = Alignment(horizontal="right")
            
            # 교대로 나타나는 행 배경색
            data_fill_light = PatternFill(start_color="e8f4f8", end_color="e8f4f8", fill_type="solid")
            data_fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            
            for row_num, product in enumerate(products, 2):
                # 교대로 배경색 설정
                current_fill = data_fill_light if (row_num - 2) % 2 == 0 else data_fill_white
                
                # ID
                id_cell = ws.cell(row=row_num, column=1)
                id_cell.value = product[0]
                id_cell.alignment = data_alignment_center
                id_cell.fill = current_fill
                
                # 제품명
                name_cell = ws.cell(row=row_num, column=2)
                name_cell.value = product[1]
                name_cell.fill = current_fill
                
                # 가격
                price_cell = ws.cell(row=row_num, column=3)
                price_cell.value = product[2]
                price_cell.alignment = data_alignment_right
                price_cell.number_format = '#,##0'
                price_cell.fill = current_fill
            
            # 컬럼 너비 설정
            ws.column_dimensions['A'].width = 10
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 20
            
            # 파일명 (현재 날짜 포함)
            filename = f"Products_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            # 파일 저장
            wb.save(filename)
            
            success_box = QMessageBox(self)
            success_box.setWindowTitle("✅ 성공")
            success_box.setText(f"엑셀 파일이 저장되었습니다.\n파일명: {filename}")
            success_box.setIcon(QMessageBox.Icon.Information)
            success_box.exec()
        
        except Exception as e:
            error_box = QMessageBox(self)
            error_box.setWindowTitle("❌ 오류")
            error_box.setText(f"엑셀 파일 저장 중 오류가 발생했습니다.\n{str(e)}")
            error_box.setIcon(QMessageBox.Icon.Critical)
            error_box.exec()


def main():
    """메인 함수"""
    app = QApplication(sys.argv)
    window = ProductsGUI()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
